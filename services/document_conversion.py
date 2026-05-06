

from __future__ import annotations

import io
import re
import logging
import tempfile
from pathlib import Path
from typing import Callable

import html2text
import pdfplumber

from bs4 import BeautifulSoup
from docx import Document
from playwright.sync_api import sync_playwright
from pdf2docx import Converter
import pypandoc

from models.enums import OutputFormat
from utils.acessories import libreoffice_convert
from utils.pdf_builder import PDFBuilder
from utils.templates import (
    clean_text,
    docx_to_html,
    epub_to_html_body,
    html_template,
    html_to_docx_bytes,
    make_epub,
)
logger = logging.getLogger(__name__)


class DocumentConversionError(Exception):
    pass


class DocumentConverter:

    _DISPATCH: dict[tuple[str, str], Callable[[bytes], bytes]] = {}

    @staticmethod
    def pdf_to_txt(content: bytes) -> bytes:
        buf = io.BytesIO(content)
        try:
            parts = []

            try:
                with pdfplumber.open(buf) as pdf:
                    for page in pdf.pages:
                        try:
                            parts.append(page.extract_text() or "")
                        except Exception:
                            parts.append("")
            except Exception as e:
                return f"Failed to extract text: {e}".encode("utf-8")

            text = "\n".join(parts).strip()
            return text.encode("utf-8") if text else b"No readable text could be extracted from this PDF."
        finally:
            buf.close()

    @staticmethod
    def pdf_to_txt_to_epub(content: bytes) -> bytes:
        try:
            buf = io.BytesIO(content)
            text_parts = []

            with pdfplumber.open(buf) as pdf:
                for page in pdf.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text_parts.append(extracted)

            text = "\n".join(text_parts).strip()

            if not text:
                text = "No readable text could be extracted from this PDF."

            return text.encode("utf-8")

        except Exception as e:
            return f"Failed to extract text: {e}".encode("utf-8")

    @staticmethod
    def pdf_to_html(content: bytes) -> bytes:
        txt = DocumentConverter.pdf_to_txt(content).decode("utf-8")

        lines = txt.splitlines()
        body_parts: list[str] = []
        for line in lines:
            s = line.strip()
            if not s:
                continue
            if s.isupper() and len(s) < 80:
                body_parts.append(f"<h2>{s}</h2>")
            else:
                body_parts.append(f"<p>{s}</p>")
        html = html_template("\n".join(body_parts))
        return html.encode("utf-8")

    @staticmethod
    def pdf_to_epub(content: bytes) -> bytes:
        raw = DocumentConverter.pdf_to_txt_to_epub(content)

        txt = raw.decode(
            "utf-8", errors="replace") if isinstance(raw, bytes) else (raw or "")

        if not txt.strip():
            txt = "No readable text could be extracted from this PDF."

        html_lines = []
        for line in txt.splitlines():
            line = line.strip()
            if not line:
                continue

            if line.isupper() and len(line) < 80:
                html_lines.append(f"<h2>{line}</h2>")
            else:
                html_lines.append(f"<p>{line}</p>")

        html_body = "\n".join(html_lines)

        return make_epub(html_body)

    @staticmethod
    def pdf_to_docx_v1(content: bytes) -> bytes:
        html = DocumentConverter.pdf_to_html(content).decode("utf-8")
        return html_to_docx_bytes(html)
    @classmethod
    def pdf_to_docx(cls, content: bytes) -> bytes:
        
       
        try:
            return cls._pdf_to_docx_pdf2docx(content)
        except Exception as e:
            logger.warning(f"[pdf2docx failed] → falling back to HTML route: {e}")

      
        try:
            html = cls.pdf_to_html(content).decode("utf-8")
            return cls.html_to_docx(html)
        except Exception as e:
            logger.error(f"[fallback failed] PDF → HTML → DOCX: {e}")
            raise DocumentConversionError("Failed to convert PDF to DOCX")
    @staticmethod
    def _pdf_to_docx_pdf2docx(content: bytes) -> bytes:
       
        with tempfile.TemporaryDirectory() as tmp:
            pdf_path = Path(tmp) / "input.pdf"
            docx_path = Path(tmp) / "output.docx"

            pdf_path.write_bytes(content)

            cv = Converter(str(pdf_path))
            try:
                cv.convert(str(docx_path))
            finally:
                cv.close()

            if not docx_path.exists():
                raise DocumentConversionError("DOCX file not generated")

            return docx_path.read_bytes()

  
    @staticmethod
    def html_to_docx(html: str) -> bytes:
        
        try:

            
            try:
                pypandoc.get_pandoc_path()
            except OSError:
                logger.info("Pandoc not found. Downloading...")
                pypandoc.download_pandoc()

            output = pypandoc.convert_text(
                html,
                to="docx",
                format="html",
                extra_args=["--quiet"]
            )

            return output

        except Exception as e:
            logger.error(f"[pypandoc failed]: {e}")
            raise DocumentConversionError("HTML → DOCX conversion failed")

    @staticmethod
    def html_to_pdf(content: bytes) -> bytes:
        html_str = content.decode("utf-8", errors="replace")

        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            page.set_content(html_str)

            pdf_bytes = page.pdf()
            browser.close()

        return pdf_bytes

    @staticmethod
    def html_to_txt(content: bytes) -> bytes:
        html_str = content.decode("utf-8", errors="replace")
        converter = html2text.HTML2Text()
        converter.ignore_links = False
        converter.body_width = 0
        return clean_text(converter.handle(html_str)).encode("utf-8")

    @staticmethod
    def html_to_epub(content: bytes) -> bytes:
        html_str = content.decode("utf-8", errors="replace")
        soup = BeautifulSoup(html_str, "html.parser")
        title_tag = soup.find("title")
        title = title_tag.get_text(strip=True) if title_tag else "Document"
        body = soup.find("body")
        body_html = body.decode_contents() if body else html_str
        return make_epub(body_html)

    @staticmethod
    def html_to_docx_v1(content: bytes) -> bytes:
        return html_to_docx_bytes(content.decode("utf-8", errors="replace"))

    @staticmethod
    def docx_to_pdf(content: bytes) -> bytes:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "input.docx"
            src.write_bytes(content)
            return libreoffice_convert(src, "pdf", Path(tmp))

    @staticmethod
    def docx_to_html(content: bytes) -> bytes:
        html_body = docx_to_html(content)
        full_html = html_template(html_body)
        return full_html.encode("utf-8")

    @staticmethod
    def docx_to_txt(content: bytes) -> bytes:
        doc = Document(io.BytesIO(content))
        lines: list[str] = []
        for para in doc.paragraphs:
            lines.append(para.text)
        return clean_text("\n".join(lines)).encode("utf-8")

    @staticmethod
    def docx_to_epub(content: bytes) -> bytes:
        html_body = docx_to_html(content)
        return make_epub(html_body)

    @staticmethod
    def txt_to_pdf(content: bytes) -> bytes:
        text = clean_text(content.decode("utf-8", errors="replace"))
        builder = PDFBuilder()
        story = builder.txt_to_story(text)
        return builder.build(story)

    @staticmethod
    def txt_to_html(content: bytes) -> bytes:
        text = clean_text(content.decode("utf-8", errors="replace"))
        body_parts: list[str] = []
        for line in text.splitlines():
            s = line.strip()
            if not s:
                body_parts.append("<br>")
                continue
            if s.startswith("# "):
                body_parts.append(f"<h1>{s[2:]}</h1>")
            elif s.startswith("## "):
                body_parts.append(f"<h2>{s[3:]}</h2>")
            elif s.startswith("### "):
                body_parts.append(f"<h3>{s[4:]}</h3>")
            elif s.startswith(("- ", "* ", "+ ")):
                body_parts.append(f"<ul><li>{s[2:]}</li></ul>")
            elif "|" in s and not re.match(r"^[\-\| :]+$", s):
                row_cells = [c.strip() for c in s.split("|") if c.strip()]
                cells_html = "".join(f"<td>{c}</td>" for c in row_cells)
                body_parts.append(f"<table><tr>{cells_html}</tr></table>")
            else:
                formatted = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", s)
                formatted = re.sub(r"\*(.*?)\*", r"<em>\1</em>", formatted)
                body_parts.append(f"<p>{formatted}</p>")
        html = html_template("\n".join(body_parts))
        return html.encode("utf-8")

    @staticmethod
    def txt_to_epub(content: bytes) -> bytes:
        html_bytes = DocumentConverter.txt_to_html(content)
        return DocumentConverter.html_to_epub(html_bytes)

    # @staticmethod
    # def txt_to_docx(content: bytes) -> bytes:
    #     text = clean_text(content.decode("utf-8", errors="replace"))
    #     doc = Document()

    #     style = doc.styles["Normal"]
    #     style.font.name = "Calibri"
    #     style.font.size = Pt(11)

    #     for line in text.splitlines():
    #         s = line.strip()

    #         if not s:
    #             doc.add_paragraph()
    #             continue

    #         if s.startswith("# "):
    #             doc.add_heading(s[2:], level=1)
    #         elif s.startswith("## "):
    #             doc.add_heading(s[3:], level=2)
    #         elif s.startswith("### "):
    #             doc.add_heading(s[4:], level=3)
    #         elif s.startswith(("- ", "* ", "+ ")):
    #             doc.add_paragraph(s[2:], style="List Bullet")
    #         elif re.match(r"^\d+\.\s", s):
    #             doc.add_paragraph(re.sub(r"^\d+\.\s", "", s), style="List Number")
    #         elif "|" in s and not re.match(r"^[\-\| :]+$", s):
    #             pass
    #         else:
    #             para = doc.add_paragraph()
    #             parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", s)
    #             for part in parts:
    #                 if part.startswith("**") and part.endswith("**"):
    #                     run = para.add_run(part[2:-2])
    #                     run.bold = True
    #                 elif part.startswith("*") and part.endswith("*"):
    #                     run = para.add_run(part[1:-1])
    #                     run.italic = True
    #                 else:
    #                     para.add_run(part)

    #     buf = io.BytesIO()
    #     doc.save(buf)
    #     buf.seek(0)
    #     return buf.read()

    @staticmethod
    def txt_to_docx(content: bytes) -> bytes:
        text = content.decode("utf-8", errors="replace")
        doc = Document()

        # title = doc.add_heading("Converted Document", level=0)
        # title.alignment = 1

        # doc.add_paragraph(
        #     f"Generated on {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
        # )

        for line in text.splitlines():
            stripped = line.strip()

            if not stripped:
                doc.add_paragraph("")
                continue

            if stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=1)
                continue

            if stripped.startswith("- ") or stripped.startswith("* "):
                doc.add_paragraph(stripped[2:], style="List Bullet")
                continue

            if stripped.isupper() and len(stripped) < 60:
                doc.add_heading(stripped, level=2)
                continue

            if "|" in stripped:
                continue

            paragraph = doc.add_paragraph()

            parts = re.split(r"(\*\*.*?\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = paragraph.add_run(part[2:-2])
                    run.bold = True
                else:
                    paragraph.add_run(part)

        buffer = io.BytesIO()
        try:
            doc.save(buffer)
            buffer.seek(0)
            data = buffer.read()
        finally:
            buffer.close()

        return data

    @staticmethod
    def epub_to_txt(content: bytes) -> bytes:
        _, body_html = epub_to_html_body(content)
        converter = html2text.HTML2Text()
        converter.ignore_links = False
        converter.body_width = 0
        return clean_text(converter.handle(body_html)).encode("utf-8")

    @staticmethod
    def epub_to_html(content: bytes) -> bytes:
        title, body_html = epub_to_html_body(content)
        return html_template(body_html).encode("utf-8")

    @staticmethod
    def epub_to_docx(content: bytes) -> bytes:
        _, body_html = epub_to_html_body(content)
        return html_to_docx_bytes(body_html)

    @staticmethod
    def epub_to_pdf(content: bytes) -> bytes:
        html_bytes = DocumentConverter.epub_to_html(content)
        return DocumentConverter.html_to_pdf(html_bytes)

    @classmethod
    def convert(cls, content: bytes, mime: str, output_format: OutputFormat) -> bytes:
        if "txt" in mime or "text/plain" in mime:
            if output_format == OutputFormat.PDF:
                return cls.txt_to_pdf(content)

            if output_format == OutputFormat.DOCX:
                return cls.txt_to_docx(content)

            if output_format == OutputFormat.HTML:
                return cls.txt_to_html(content)
            # if output_format == OutputFormat.EPUB:
            #     return cls.txt_to_epub(content)

        if "pdf" in mime or "application/pdf" in mime:
            if output_format == OutputFormat.TXT:
                return cls.pdf_to_txt(content)
            # if output_format == OutputFormat.EPUB:
            #     return cls.pdf_to_epub(content)
            if output_format == OutputFormat.HTML:
                return cls.pdf_to_html(content)
            if output_format == OutputFormat.DOCX:
                return cls.pdf_to_docx(content)

        if "wordprocessingml" in mime or "msword" in mime:
            if output_format == OutputFormat.TXT:
                return cls.docx_to_txt(content)

            if output_format == OutputFormat.PDF:
                return cls.docx_to_pdf(content)
            if output_format == OutputFormat.HTML:
                return cls.docx_to_html(content)
            # if output_format == OutputFormat.EPUB:
            #     return cls.docx_to_epub(content)
        if "application/epub+zip" in mime:
            if output_format == OutputFormat.TXT:
                return cls.epub_to_txt(content)

            if output_format == OutputFormat.PDF:
                return cls.epub_to_pdf(content)
            if output_format == OutputFormat.HTML:
                return cls.epub_to_html(content)
            if output_format == OutputFormat.DOCX:
                return cls.epub_to_docx(content)
        if "html" in mime or "text/html" in mime:
            if output_format == OutputFormat.TXT:
                return cls.html_to_txt(content)
            if output_format == OutputFormat.PDF:
                return cls.html_to_pdf(content)
            # if output_format == OutputFormat.EPUB:
            #     return cls.html_to_epub(content)
            if output_format == OutputFormat.DOCX:
                return cls.html_to_docx(content)

        raise ValueError(f"Unsupported: {mime} → {output_format}")
